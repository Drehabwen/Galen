# -*- coding: utf-8 -*-
"""Generate Galen beta-test feedback template (docx)."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = (0x1b, 0x2e, 0x1f)      # 墨绿（品牌主色）
ORANGE = (0xd4, 0x74, 0x3c)   # 湖橙（品牌副色）
GRAY = (0x80, 0x80, 0x80)

FONT = '微软雅黑'

def set_run(run, size=10.5, bold=False, color=None, italic=False, font=FONT):
    run.font.name = font
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run

def para(doc, text='', size=10.5, bold=False, color=None, italic=False,
         align=None, space_before=0, space_after=4, indent=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    if text:
        set_run(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p

def section_header(doc, num, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    set_run(p.add_run(f'{num}｜{title}'), size=12, bold=True, color=INK)
    # 下边框线
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '2'); bottom.set(qn('w:color'), 'd4743c')
    pBdr.append(bottom); pPr.append(pBdr)

def blank_box(doc, lines=3, hint=None):
    """填写区：带边框的空表格，方便手写/打字。"""
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    cell = t.rows[0].cells[0]
    cell.width = Cm(16.5)
    first = True
    for i in range(lines):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(10)
        if hint and i == 0:
            set_run(p.add_run(hint), size=9, color=GRAY, italic=True)
    return t

def fill_table(doc, rows_data, widths=(4.5, 12.0)):
    t = doc.add_table(rows=len(rows_data), cols=2)
    t.style = 'Table Grid'
    for i, (label, value) in enumerate(rows_data):
        c0, c1 = t.rows[i].cells
        c0.width = Cm(widths[0]); c1.width = Cm(widths[1])
        p0 = c0.paragraphs[0]; set_run(p0.add_run(label), size=10, bold=True)
        p1 = c1.paragraphs[0]; set_run(p1.add_run(value), size=10, color=GRAY if value.startswith('（') else None)
    return t

doc = Document()
# 页边距
for s in doc.sections:
    s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)

# ============ 标题区 ============
para(doc, 'Galen', size=22, bold=True, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
para(doc, '医学科研助手 · 内测用户反馈表', size=14, bold=True, color=ORANGE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
para(doc, '感谢参与 Galen 内测！你的每一条反馈都会直接进入开发流程。\n请尽量完整填写，尤其是「复现步骤」——它决定了问题能否被快速定位与修复。',
     size=9.5, color=GRAY, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

# ============ 一、基本信息 ============
section_header(doc, '01', '基本信息')
fill_table(doc, [
    ('反馈人（姓名/昵称）', ''),
    ('联系方式（微信/邮箱，选填）', ''),
    ('反馈日期', ''),
    ('Galen 版本', 'v0.2.0（主界面左下角 / 关于页可查看）'),
    ('操作系统', '□ Windows 10   □ Windows 11   版本号：____________'),
    ('设备内存', '□ 8GB   □ 16GB   □ 32GB 及以上'),
])

# ============ 二、问题类型 ============
section_header(doc, '02', '问题类型（可多选）')
types = [
    '□ 程序崩溃 / 闪退', '□ 功能异常（功能不工作、结果不对）', '□ 界面显示问题（布局/字体/错位）',
    '□ 性能问题（卡顿 / 启动慢 / 占用高）', '□ 使用困惑（不知道怎么操作）', '□ 功能建议 / 新想法',
    '□ 其他：________________________________',
]
for i in range(0, len(types), 2):
    line = '        '.join(types[i:i+2])
    para(doc, line, size=10, space_after=3)

# ============ 三、问题描述 ============
section_header(doc, '03', '问题描述')
para(doc, '3.1 发生了什么？（一两句话概括）', size=10, bold=True, space_before=4)
blank_box(doc, 3)
para(doc, '3.2 复现步骤（最重要的一节！从打开 Galen 开始，按顺序写出每一步）', size=10, bold=True, space_before=8)
blank_box(doc, 5, hint='例：1. 打开 Galen → 2. 新建任务并输入「帮我分析这份量表数据」 → 3. 点击确认计划 → 4. 出现报错…')
para(doc, '3.3 期望结果 与 实际结果', size=10, bold=True, space_before=8)
fill_table(doc, [('我期望的结果', ''), ('实际发生的结果', '')], widths=(4.5, 12.0))
para(doc, '3.4 出现频率', size=10, bold=True, space_before=8)
para(doc, '□ 每次必现        □ 经常出现        □ 偶尔出现        □ 只出现过一次', size=10, space_after=3)

# ============ 四、证据材料 ============
section_header(doc, '04', '证据材料（如有请附上）')
for it in [
    '□ 截图（文件名：__________________________）',
    '□ 屏幕录像',
    '□ 报错弹窗截图（有报错信息时，请务必截图）',
    '□ 日志文件（如能找到：%USERPROFILE%\\.galen\\ 目录下的相关文件）',
    '□ 其他：______________________________',
]:
    para(doc, it, size=10, space_after=3)

# ============ 五、严重程度 ============
section_header(doc, '05', '严重程度自评')
for it in [
    '□ P0 阻塞 —— 完全无法使用',
    '□ P1 严重 —— 核心功能不可用',
    '□ P2 一般 —— 功能可用但存在明显问题',
    '□ P3 轻微 —— 不影响使用的小问题',
    '□ P4 建议 —— 改进类想法 / 体验优化',
]:
    para(doc, it, size=10, space_after=3)

# ============ 六、其他 ============
section_header(doc, '06', '其他想说的')
para(doc, '（任何使用体验、想要的新功能、其他意见……都欢迎写在这里）', size=9, color=GRAY, italic=True, space_after=4)
blank_box(doc, 5)

# ============ 自查清单 ============
section_header(doc, '✓', '提交前快速自查')
para(doc, '□ 版本号已填写      □ 复现步骤已写出      □ 截图为最新      □ 问题类型已勾选      □ 严重程度已自评',
     size=10, space_after=4)
para(doc, '提交方式：______________________________（请团队自行填写反馈接收渠道）',
     size=9.5, color=GRAY, space_before=8)

# 页脚
footer = doc.sections[0].footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(fp.add_run('Galen · 医学科研助手 —— 闭环工具：采集 · 处理 · 分析 · 成文 · 签核'), size=8, color=GRAY)

out = r'D:\Users\DORAT\Desktop\03_AI康复项目\Galen\docs\Galen_内测反馈模板.docx'
doc.save(out)
print('saved:', out)
